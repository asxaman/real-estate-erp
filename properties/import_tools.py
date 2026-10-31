"""Excel / CSV import for properties and units (owners are in owner_tools.py)."""
import re
from decimal import Decimal, InvalidOperation

from django.core.exceptions import ValidationError
from django.http import HttpResponse

from .models import Owner, Property, Unit
from .models_constants import normalise_unit_number
from .owner_tools import MAX_IMPORT_ROWS, _read_rows


class RowError(Exception):
    """A problem with one cell or row, in plain words."""


# ---- small helpers -----------------------------------------------------------------------
def _norm(text):
    return re.sub(r"[^a-z0-9]+", " ", (text or "").lower()).strip()


def _squash(text):
    return re.sub(r"[^a-z0-9]", "", (text or "").lower())


def _pick(text, choices, field):
    """Match a typed word to a choice by its stored value or its label (capitals/spaces ignored)."""
    if not text:
        return ""
    wanted = _squash(text)
    for value, label in choices:
        if wanted in (_squash(value), _squash(label)):
            return value
    allowed = ", ".join(label for _, label in choices)
    raise RowError(f'{field} "{text}" is not valid. Use one of: {allowed}.')


def _number(text, field, integer=False):
    if not text:
        return None
    cleaned = text.replace(",", "").replace(" ", "")
    try:
        value = Decimal(cleaned)
    except InvalidOperation:
        raise RowError(f'{field} "{text}" is not a number.')
    if integer:
        if value != value.to_integral_value():
            raise RowError(f'{field} "{text}" must be a whole number.')
        return int(value)
    return value


def _date(text, field):
    """2024-05-31, 31/05/2024 or an Excel date (2024-05-31 00:00:00)."""
    from datetime import datetime

    if not text:
        return None
    word = text.strip()[:10] if re.match(r"\d{4}-\d{2}-\d{2}", text.strip()) else text.strip()
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(word, pattern).date()
        except ValueError:
            continue
    raise RowError(f'{field} "{text}" is not a date. Use 2024-05-31 or 31/05/2024.')


def _yes_no(text, field):
    if not text:
        return False
    word = text.strip().lower()
    if word in ("yes", "y", "true", "1"):
        return True
    if word in ("no", "n", "false", "0"):
        return False
    raise RowError(f'{field} "{text}" must be Yes or No.')


def _describe(exc):
    if isinstance(exc, RowError):
        return str(exc)
    messages = getattr(exc, "message_dict", None)
    if messages:
        return "; ".join(f"{field}: {' '.join(msgs)}" for field, msgs in messages.items())
    return "; ".join(getattr(exc, "messages", [str(exc)]))


def _load(upload, headers, aliases):
    """Read the file and map its columns. Returns (columns, body_rows, errors)."""
    try:
        rows = _read_rows(upload)
    except ValueError as exc:
        return None, None, [str(exc)]
    except Exception:
        return None, None, ["The file could not be read. Use the template from this page."]
    if not rows:
        return None, None, ["The file is empty."]
    columns = {}
    for index, header in enumerate(rows[0]):
        key = aliases.get(_norm(header))
        if key:
            columns[key] = index
    missing = [h for h, key in headers if key.startswith("*") and key[1:] not in columns]
    if missing:
        return None, None, [
            f'The first row must contain the column(s): {", ".join(missing)}. Use the template from this page.'
        ]
    body = [(n, r) for n, r in enumerate(rows[1:], start=2) if any(c for c in r)]
    if len(body) > MAX_IMPORT_ROWS:
        return None, None, [f"The file has {len(body)} rows. The maximum is {MAX_IMPORT_ROWS} per import."]
    return columns, body, []


def _getter(columns):
    def get(row, key):
        index = columns.get(key)
        return row[index].strip() if index is not None and index < len(row) else ""

    return get


# ---- properties ----------------------------------------------------------------------------
PROPERTY_HEADERS = [
    "Code", "Name", "Type", "Owner", "Emirate", "Address", "Map link", "Year built",
    "Total floors", "Apartment type", "Bathrooms", "Annual rent (AED)", "Status",
    "Fee type", "Fee percent", "Fee fixed (AED)", "Notes",
    "Area / community", "Parking spaces", "Title deed no.", "Plot no.", "Makani no.", "Utility premise no.",
]
_PROPERTY_ALIASES = {
    "code": "code", "property code": "code",
    "name": "name", "property name": "name",
    "type": "type", "property type": "type",
    "owner": "owner", "owner name": "owner",
    "emirate": "emirate",
    "address": "address",
    "map link": "map_link", "map": "map_link",
    "year built": "year_built", "year": "year_built",
    "total floors": "total_floors", "floors": "total_floors",
    "apartment type": "apartment_type",
    "bathrooms": "bathrooms",
    "annual rent aed": "annual_rent", "annual rent": "annual_rent", "rent": "annual_rent",
    "status": "status",
    "fee type": "fee_type", "management fee type": "fee_type",
    "fee percent": "fee_percent", "fee": "fee_percent", "management fee": "fee_percent",
    "fee fixed aed": "fee_fixed", "fee fixed": "fee_fixed",
    "notes": "notes", "note": "notes",
    "area community": "community", "community": "community", "area": "community",
    "parking spaces": "parking_spaces", "parking": "parking_spaces",
    "title deed no": "title_deed", "title deed": "title_deed", "title deed number": "title_deed",
    "plot no": "plot", "plot": "plot", "plot number": "plot",
    "makani no": "makani", "makani": "makani", "makani number": "makani",
    "utility premise no": "premise", "premise no": "premise", "dewa premise no": "premise",
}
_PROPERTY_REQUIRED = [("Code", "*code"), ("Name", "*name"), ("Type", "*type")]
_PROPERTY_REQUIRED_WHEN_UPDATING = [("Code", "*code")]


def _property_values(get, row, owners, company, choices):
    """Only the cells that are filled in, turned into field values (blank cells are left out)."""
    values = {}
    owner_name = get(row, "owner")
    if owner_name:
        if owner_name.lower() == company.name.lower():
            values["owner"] = company
        else:
            owner = owners.get(owner_name.lower())
            if owner is None:
                raise RowError(
                    f'Owner "{owner_name}" does not exist. Create the owner first (or leave blank for Company Owned).'
                )
            values["owner"] = owner
    for key, attr, kind in (
        ("name", "name", "text"), ("address", "address", "text"), ("map_link", "map_link", "text"),
        ("notes", "notes", "text"), ("community", "community", "text"),
        ("title_deed", "title_deed_number", "text"), ("plot", "plot_number", "text"),
        ("makani", "makani_number", "text"), ("premise", "dewa_premise_number", "text"),
    ):
        if get(row, key):
            values[attr] = get(row, key)
    for key, attr, label in (
        ("year_built", "year_built", "Year built"), ("total_floors", "total_floors", "Total floors"),
        ("bathrooms", "bathrooms", "Bathrooms"), ("parking_spaces", "parking_spaces", "Parking spaces"),
    ):
        if get(row, key):
            values[attr] = _number(get(row, key), label, integer=True)
    if get(row, "annual_rent"):
        values["annual_rent"] = _number(get(row, "annual_rent"), "Annual rent (AED)")
    for key, attr, label, options in (
        ("type", "property_type", "Type", choices["type"]),
        ("emirate", "emirate", "Emirate", choices["emirate"]),
        ("status", "status", "Status", choices["status"]),
        ("apartment_type", "apartment_type", "Apartment type", choices["apartment"]),
    ):
        if get(row, key):
            values[attr] = _pick(get(row, key), options, label)

    fee_text = get(row, "fee_type").lower()
    fee_percent = _number(get(row, "fee_percent"), "Fee percent")
    fee_fixed = _number(get(row, "fee_fixed"), "Fee fixed (AED)")
    if fee_text or fee_percent is not None or fee_fixed is not None:
        if fee_text.startswith(("perc", "%")):
            fee_type = Property.FeeType.PERCENT
        elif fee_text.startswith("fix"):
            fee_type = Property.FeeType.FIXED
        elif fee_text:
            raise RowError(f'Fee type "{get(row, "fee_type")}" is not valid. Use Percent or Fixed.')
        elif fee_fixed:
            fee_type = Property.FeeType.FIXED
        elif fee_percent:
            fee_type = Property.FeeType.PERCENT
        else:
            fee_type = ""
        values["management_fee_type"] = fee_type
        values["management_fee_percent"] = fee_percent if fee_percent is not None else 0
        values["management_fee_fixed"] = fee_fixed
    return values


def parse_property_file(upload, update=False):
    """
    Returns (properties, errors). Nothing is saved; the caller saves when there are no errors.
    With update=True a row whose Code already exists changes that property (blank cells are left alone).
    """
    required = _PROPERTY_REQUIRED_WHEN_UPDATING if update else _PROPERTY_REQUIRED
    columns, body, errors = _load(upload, required, _PROPERTY_ALIASES)
    if errors:
        return [], errors
    get = _getter(columns)
    owners = {o.name.lower(): o for o in Owner.objects.all()}
    company = Owner.company()
    choices = {
        "type": Property.PropertyType.choices,
        "emirate": Property.Emirate.choices,
        "status": Property.Status.choices,
        "apartment": Property._meta.get_field("apartment_type").choices,
    }
    existing = {p.code.upper(): p for p in Property.objects.select_related("owner")} if update else {}

    result, seen = [], set()
    for line, row in body:
        try:
            code = get(row, "code").upper()
            values = _property_values(get, row, owners, company, choices)
            current = existing.get(code)
            if current is not None:
                prop = current
                for attr, value in values.items():
                    setattr(prop, attr, value)
            else:
                if not values.get("property_type"):
                    raise RowError("Type is required.")
                fields = dict(
                    owner=company, emirate=Property.Emirate.ABU_DHABI, management_fee_type="",
                    management_fee_percent=0,
                )
                fields.update(values)
                prop = Property(code=code, **fields)
                if not prop.has_units and not prop.status:
                    prop.status = Property.Status.VACANT
            prop.full_clean(exclude=["created_at"])
            if prop.code in seen:
                raise RowError(f'Code "{prop.code}" appears twice in the file.')
            seen.add(prop.code)
            result.append(prop)
        except (RowError, ValidationError) as exc:
            errors.append(f"Row {line}: {_describe(exc)}")
    return result, errors


# ---- units -----------------------------------------------------------------------------------
UNIT_HEADERS = [
    "Property code", "Unit number", "Floor", "Type", "Bathrooms", "Area (sq ft)",
    "Furnished", "Status", "Asking rent (AED / year)", "Notes",
    "Parking spaces", "View", "Utility account no.",
]
_UNIT_ALIASES = {
    "property code": "property", "property": "property", "code": "property", "building": "property",
    "unit number": "number", "unit": "number", "number": "number", "unit no": "number",
    "floor": "floor",
    "type": "type", "unit type": "type",
    "bathrooms": "bathrooms",
    "area sq ft": "area", "area": "area",
    "furnished": "furnished",
    "status": "status",
    "asking rent aed year": "rent", "asking rent": "rent", "rent": "rent",
    "notes": "notes", "note": "notes",
    "parking spaces": "parking_spaces", "parking": "parking_spaces",
    "view": "view",
    "utility account no": "account", "dewa account no": "account", "account no": "account",
}
_UNIT_REQUIRED = [("Property code", "*property"), ("Unit number", "*number"), ("Type", "*type")]
_UNIT_REQUIRED_WHEN_UPDATING = [("Property code", "*property"), ("Unit number", "*number")]


def parse_unit_file(upload, update=False):
    """Like parse_property_file: with update=True a row for an existing unit changes only what it fills in."""
    required = _UNIT_REQUIRED_WHEN_UPDATING if update else _UNIT_REQUIRED
    columns, body, errors = _load(upload, required, _UNIT_ALIASES)
    if errors:
        return [], errors
    get = _getter(columns)
    properties = {p.code.upper(): p for p in Property.objects.all()}
    type_choices = Unit.UnitType.choices
    status_choices = Unit.Status.choices
    view_choices = Unit.View.choices
    existing = {(u.property_id, u.unit_number): u for u in Unit.objects.select_related("property")} if update else {}

    result, seen = [], set()
    for line, row in body:
        try:
            code = get(row, "property").upper()
            prop = properties.get(code)
            if prop is None:
                raise RowError(f'Property code "{get(row, "property")}" does not exist.')
            number = normalise_unit_number(get(row, "number"))
            values = {}
            if get(row, "floor"):
                values["floor"] = get(row, "floor").upper()
            if get(row, "type"):
                values["unit_type"] = _pick(get(row, "type"), type_choices, "Type")
            if get(row, "bathrooms"):
                values["bathrooms"] = _number(get(row, "bathrooms"), "Bathrooms", integer=True)
            if get(row, "area"):
                values["area_sqft"] = _number(get(row, "area"), "Area (sq ft)")
            if get(row, "furnished"):
                values["furnished"] = _yes_no(get(row, "furnished"), "Furnished")
            if get(row, "status"):
                values["status"] = _pick(get(row, "status"), status_choices, "Status")
            if get(row, "rent"):
                values["asking_rent"] = _number(get(row, "rent"), "Asking rent")
            if get(row, "notes"):
                values["notes"] = get(row, "notes")
            if get(row, "parking_spaces"):
                values["parking_spaces"] = _number(get(row, "parking_spaces"), "Parking spaces", integer=True)
            if get(row, "view"):
                values["view"] = _pick(get(row, "view"), view_choices, "View")
            if get(row, "account"):
                values["dewa_account_number"] = get(row, "account")

            current = existing.get((prop.pk, number))
            if current is not None:
                unit = current
                for attr, value in values.items():
                    setattr(unit, attr, value)
            else:
                if not values.get("unit_type"):
                    raise RowError("Type is required.")
                fields = dict(floor="", furnished=False, status=Unit.Status.VACANT)
                fields.update(values)
                unit = Unit(property=prop, unit_number=number, **fields)
            unit.full_clean(exclude=["floor_sort", "created_at"])
            key = (prop.pk, unit.unit_number)
            if key in seen:
                raise RowError(f'Unit "{code}-{unit.unit_number}" appears twice in the file.')
            seen.add(key)
            result.append(unit)
        except (RowError, ValidationError) as exc:
            errors.append(f"Row {line}: {_describe(exc)}")
    return result, errors


# ---- templates -------------------------------------------------------------------------------
def _template(title, headers, example, filename, allowed):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.append(headers)
    sheet.append(example)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for index in range(1, len(headers) + 1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = max(14, len(headers[index - 1]) + 4)
    help_sheet = workbook.create_sheet("Allowed values")
    for name, values in allowed:
        help_sheet.append([name, ", ".join(values)])
    help_sheet.column_dimensions["A"].width = 18
    help_sheet.column_dimensions["B"].width = 100
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    workbook.save(response)
    return response


def property_template_response():
    return _template(
        "Properties", PROPERTY_HEADERS,
        ["RMT", "Rimal Tower", "Residential Building", "", "Abu Dhabi", "Corniche Road", "", 2015, 12,
         "", "", "", "", "", "", "", "", "", "", "", "", "", ""],
        "properties-import-template.xlsx",
        [
            ("Type", [label for _, label in Property.PropertyType.choices]),
            ("Emirate", [label for _, label in Property.Emirate.choices]),
            ("Status", [label for _, label in Property.Status.choices] + ["(standalone only; blank = Vacant)"]),
            ("Fee type", ["Percent", "Fixed", "(not needed for Company Owned)"]),
            ("Owner", ["Exact owner name; blank = Company Owned"]),
        ],
    )


def unit_template_response():
    return _template(
        "Units", UNIT_HEADERS,
        ["RMT", "101", "1", "2BHK", 3, 1200, "No", "Vacant", 95000, "", 1, "Sea", ""],
        "units-import-template.xlsx",
        [
            ("Type", [label for _, label in Unit.UnitType.choices]),
            ("Status", [label for _, label in Unit.Status.choices]),
            ("Furnished", ["Yes", "No"]),
            ("Floor", ["G, B1, P1, M, 1, 2, 3 ... (as in the unit form)"]),
        ],
    )
