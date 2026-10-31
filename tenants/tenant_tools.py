"""Excel / CSV export and import for tenants."""
import csv
from datetime import date

from django.core.exceptions import ValidationError
from django.http import HttpResponse

from core import countries
from core.exports import safe_cell
from properties.import_tools import RowError, _date, _describe, _getter, _load, _pick
from properties.models_constants import lost_zero_problem

from .models import Tenant, same_person

EXPORT_HEADERS = [
    "Code", "Name", "Type", "Active", "Blacklisted", "Nationality", "Email", "Phone", "Address",
    "Emirates ID", "Passport", "Date of birth", "Trade licence", "Authorised signatory", "TRN",
    "Documents", "Notes",
]
IMPORT_HEADERS = [
    "Name", "Type", "Nationality", "Email", "Phone", "Address", "Notes",
    "Emirates ID", "Passport", "Date of birth", "Trade licence", "Authorised signatory", "TRN",
]
_REQUIRED = [("Name", "*name"), ("Nationality", "*nationality"), ("Email", "*email"), ("Phone", "*phone")]
_REQUIRED_WHEN_UPDATING = [("Email", "*email")]
_ALIASES = {
    "name": "name", "tenant name": "name",
    "type": "type", "tenant type": "type",
    "nationality": "nationality",
    "email": "email", "e mail": "email",
    "phone": "phone", "mobile": "phone", "telephone": "phone",
    "address": "address",
    "notes": "notes", "note": "notes",
    "emirates id": "eid", "emirates id no": "eid", "emirates id number": "eid",
    "passport": "passport", "passport no": "passport", "passport number": "passport",
    "date of birth": "dob", "birth date": "dob", "dob": "dob",
    "trade licence": "licence", "trade license": "licence", "trade licence no": "licence",
    "authorised signatory": "signatory", "authorized signatory": "signatory", "signatory": "signatory",
    "trn": "trn", "tax registration number": "trn",
}
_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _docs_label(tenant):
    expired = getattr(tenant, "_docs_expired", None)
    soon = getattr(tenant, "_docs_soon", None)
    if expired is None:
        expired, soon = tenant.document_counts()
    if expired:
        return f"{expired} expired"
    if soon:
        return f"{soon} expiring soon"
    return "OK"


def export_rows(queryset):
    for t in queryset:
        yield [
            t.code or "", t.name, t.get_tenant_type_display(), "Yes" if t.is_active else "No",
            "Yes" if t.is_blacklisted else "No", t.nationality, t.email, t.phone, t.address,
            t.emirates_id_number, t.passport_number, t.date_of_birth.isoformat() if t.date_of_birth else "",
            t.trade_licence_number, t.authorised_signatory, t.trn, _docs_label(t), t.notes,
        ]


def _filename(extension):
    return f"tenants-{date.today().isoformat()}.{extension}"


def excel_response(queryset):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tenants"
    sheet.append(EXPORT_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in export_rows(queryset):
        sheet.append([safe_cell(value) for value in row])
    sheet.freeze_panes = "A2"
    for column, width in zip("ABCDEFGHIJKLMNOPQ", [11, 30, 12, 8, 11, 20, 28, 18, 34, 20, 14, 14, 18, 24, 18, 16, 36]):
        sheet.column_dimensions[column].width = width
    response = HttpResponse(content_type=_XLSX)
    response["Content-Disposition"] = f'attachment; filename="{_filename("xlsx")}"'
    workbook.save(response)
    return response


def csv_response(queryset):
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{_filename("csv")}"'
    response.write("﻿")  # so Excel reads Arabic names correctly
    writer = csv.writer(response)
    writer.writerow(EXPORT_HEADERS)
    for row in export_rows(queryset):
        writer.writerow([safe_cell(value) for value in row])
    return response


def template_response():
    from openpyxl import Workbook
    from openpyxl.styles import Font

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tenants"
    sheet.append(IMPORT_HEADERS)
    sheet.append(["Example Tenant", "Individual", "Egypt", "tenant@example.com", "+971 50 123 4567", "Abu Dhabi", "",
                  "784-1990-1234567-1", "A1234567", "1990-05-31", "", "", ""])
    for row in range(2, 1002):  # Phone column as Text, so Excel keeps a leading 0 or +
        sheet.cell(row=row, column=5).number_format = "@"
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for column, width in zip("ABCDEFGHIJKLM", [30, 12, 16, 28, 18, 34, 36, 22, 14, 14, 16, 22, 18]):
        sheet.column_dimensions[column].width = width
    response = HttpResponse(content_type=_XLSX)
    response["Content-Disposition"] = 'attachment; filename="tenants-import-template.xlsx"'
    workbook.save(response)
    return response


def _country(text):
    """'Egyptian' / 'egypt' / 'UAE' -> the country name used in the dropdown."""
    if not text:
        return ""
    country = countries.match(text)
    if not country:
        raise RowError(
            f'Nationality "{text}" is not recognised. Use a country name such as Egypt or United Kingdom.'
        )
    return country


def parse_tenant_file(upload, update=False):
    """
    Returns (tenants, errors). Nothing is saved; the caller saves when there are no errors.
    With update=True a row whose Email matches an existing tenant changes that tenant
    (blank cells are left alone).
    """
    columns, body, errors = _load(upload, _REQUIRED_WHEN_UPDATING if update else _REQUIRED, _ALIASES)
    if errors:
        return [], errors
    get = _getter(columns)
    types = Tenant.TenantType.choices
    existing = {t.email.lower(): t for t in Tenant.objects.all() if t.email} if update else {}
    result, seen = [], []
    for line, row in body:
        try:
            problem = lost_zero_problem(get(row, "phone"))
            if problem:
                raise RowError(f'Phone "{get(row, "phone")}" {problem}')
            values = {}
            for key, attr in (("name", "name"), ("email", "email"), ("phone", "phone"), ("address", "address"),
                              ("notes", "notes"), ("eid", "emirates_id_number"), ("passport", "passport_number"),
                              ("licence", "trade_licence_number"), ("signatory", "authorised_signatory"),
                              ("trn", "trn")):
                if get(row, key):
                    values[attr] = get(row, key)
            if get(row, "type"):
                values["tenant_type"] = _pick(get(row, "type"), types, "Type")
            if get(row, "nationality"):
                values["nationality"] = _country(get(row, "nationality"))
            if get(row, "dob"):
                values["date_of_birth"] = _date(get(row, "dob"), "Date of birth")

            current = existing.get(get(row, "email").lower())
            if current is not None:
                tenant = current
                for attr, value in values.items():
                    setattr(tenant, attr, value)
            else:
                values.setdefault("tenant_type", Tenant.TenantType.INDIVIDUAL)
                values.setdefault("nationality", "")
                tenant = Tenant(**values)
            tenant.full_clean(exclude=["photo", "code"])
            person = (tenant.name, tenant.email, tenant.phone)
            if any(same_person(person, earlier) for earlier in seen):
                raise RowError(
                    f'"{tenant.name}" looks like a duplicate of an earlier row '
                    "(same email, or same name and phone number)."
                )
            seen.append(person)
            result.append(tenant)
        except (RowError, ValidationError) as exc:
            errors.append(f"Row {line}: {_describe(exc)}")
    return result, errors
